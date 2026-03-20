from pathlib import Path
from unittest import mock

import pytest

from openpyxl import Workbook

from excel_personal_data_validator.cli import main
from excel_personal_data_validator.db import NameCategory, NameDatabase


def test_main_all_valid(tmp_path: Path, capsys):
    """Тест: все данные корректны."""
    xlsx = tmp_path / "data.xlsx"
    db_path = tmp_path / "names.db"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["Иванов", "Иван", "Петрович"])
    wb.save(xlsx)
    wb.close()

    # Предварительно наполняем БД
    with NameDatabase(db_path) as db:
        db.initialize()
        db.add_name(NameCategory.LAST_NAME, "Иванов")
        db.add_name(NameCategory.FIRST_NAME, "Иван")
        db.add_name(NameCategory.PATRONYMIC, "Петрович")

    with mock.patch("sys.argv", ["validator", str(xlsx), "--db-path", str(db_path)]):
        main()

    captured = capsys.readouterr()
    assert "корректны" in captured.out


def test_main_with_unknowns_skip(tmp_path: Path, capsys):
    """Тест: есть неизвестные, пользователь пропускает."""
    xlsx = tmp_path / "data.xlsx"
    db_path = tmp_path / "names.db"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["НеизвФам", "НеизвИмя", "НеизвОтч"])
    wb.save(xlsx)
    wb.close()

    with NameDatabase(db_path) as db:
        db.initialize()

    # Пропускаем все 3 неизвестных (для каждого: [1] Ввести, [2] Добавить, [3] Пропустить → 3)
    with (
        mock.patch("sys.argv", ["validator", str(xlsx), "--db-path", str(db_path)]),
        mock.patch("builtins.input", return_value="3"),
    ):
        main()

    captured = capsys.readouterr()
    assert "Пропущено" in captured.out


def test_main_with_corrections(tmp_path: Path, capsys):
    """Тест: пользователь добавляет значения в БД."""
    xlsx = tmp_path / "data.xlsx"
    db_path = tmp_path / "names.db"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["Новиков", "Новое", "Новоич"])
    wb.save(xlsx)
    wb.close()

    with NameDatabase(db_path) as db:
        db.initialize()

    # Добавляем все 3 в базу (для каждого: [1] Ввести, [2] Добавить, [3] Пропустить → 2)
    with (
        mock.patch("sys.argv", ["validator", str(xlsx), "--db-path", str(db_path)]),
        mock.patch("builtins.input", return_value="2"),
    ):
        main()

    captured = capsys.readouterr()
    assert "Добавлено в базу" in captured.out

    # Проверяем что данные в БД
    with NameDatabase(db_path) as db:
        assert "новиков" in db.load_all(NameCategory.LAST_NAME)


def test_main_empty_file(tmp_path: Path, capsys):
    """Тест: пустой файл."""
    xlsx = tmp_path / "empty.xlsx"
    db_path = tmp_path / "names.db"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    wb.save(xlsx)
    wb.close()

    with NameDatabase(db_path) as db:
        db.initialize()

    with mock.patch("sys.argv", ["validator", str(xlsx), "--db-path", str(db_path)]):
        main()

    captured = capsys.readouterr()
    assert "пуст" in captured.out.lower() or "данных" in captured.out.lower()


def test_import_db_first_names(tmp_path: Path, capsys):
    """Тест: импорт имён из файла."""
    db_path = tmp_path / "names.db"
    names_file = tmp_path / "first_names.txt"
    names_file.write_text("Иван\nПётр\nМария\n", encoding="utf-8")

    with mock.patch(
        "sys.argv", ["validator", "import-db", "--first-names", str(names_file), "--db-path", str(db_path)]
    ):
        main()

    captured = capsys.readouterr()
    assert "Имя: 3" in captured.out

    with NameDatabase(db_path) as db:
        names = db.load_all(NameCategory.FIRST_NAME)
        assert "иван" in names
        assert "пётр" in names
        assert "мария" in names


def test_import_db_all_categories(tmp_path: Path, capsys):
    """Тест: импорт всех категорий."""
    db_path = tmp_path / "names.db"

    last = tmp_path / "last.txt"
    last.write_text("Иванов\nПетров\n", encoding="utf-8")
    first = tmp_path / "first.txt"
    first.write_text("Иван\n", encoding="utf-8")
    patr = tmp_path / "patr.txt"
    patr.write_text("Иванович\nПетрович\nСергеевич\n", encoding="utf-8")

    with mock.patch(
        "sys.argv",
        [
            "validator",
            "import-db",
            "--last-names",
            str(last),
            "--first-names",
            str(first),
            "--patronymics",
            str(patr),
            "--db-path",
            str(db_path),
        ],
    ):
        main()

    captured = capsys.readouterr()
    assert "Фамилия: 2" in captured.out
    assert "Имя: 1" in captured.out
    assert "Отчество: 3" in captured.out
    assert "Всего: 6" in captured.out


def test_import_db_no_files(tmp_path: Path, capsys):
    """Тест: вызов import-db без файлов."""
    db_path = tmp_path / "names.db"

    with mock.patch("sys.argv", ["validator", "import-db", "--db-path", str(db_path)]), pytest.raises(SystemExit):
        main()

    captured = capsys.readouterr()
    assert "укажите" in captured.out.lower() or "Ошибка" in captured.out


def test_import_db_file_not_found(tmp_path: Path, capsys):
    """Тест: файл для импорта не найден."""
    db_path = tmp_path / "names.db"

    with (
        mock.patch(
            "sys.argv",
            ["validator", "import-db", "--first-names", str(tmp_path / "nonexistent.txt"), "--db-path", str(db_path)],
        ),
        pytest.raises(SystemExit),
    ):
        main()

    captured = capsys.readouterr()
    assert "не найден" in captured.out
