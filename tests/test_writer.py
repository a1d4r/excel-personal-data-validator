from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_personal_data_validator.writer import save_corrected_excel


def test_save_corrected_excel(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["Ивано", "Иван", "Петрович"])
    ws.append(["Петров", "Алексй", "Сергеевич"])
    wb.save(source)
    wb.close()

    corrections = {(2, "A"): "Иванов", (3, "B"): "Алексей"}
    save_corrected_excel(source, output, corrections)

    wb = load_workbook(output)
    ws = wb.active
    assert ws is not None
    assert ws["A2"].value == "Иванов"
    assert ws["B3"].value == "Алексей"
    # Неизменённые ячейки
    assert ws["B2"].value == "Иван"
    assert ws["C2"].value == "Петрович"
    wb.close()


def test_save_corrected_excel_no_corrections(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Данные"])
    wb.save(source)
    wb.close()

    save_corrected_excel(source, output, {})
    assert not output.exists()
