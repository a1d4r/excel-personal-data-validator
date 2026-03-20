from pathlib import Path

from openpyxl import Workbook

from excel_personal_data_validator.config import AppConfig
from excel_personal_data_validator.reader import get_output_path, read_excel


def test_read_excel_basic(sample_xlsx: Path):
    config = AppConfig(excel_path=sample_xlsx, db_path=Path("unused.db"), start_row=2)
    rows = read_excel(config)
    assert len(rows) == 4
    assert rows[0].last_name == "Иванов"
    assert rows[0].first_name == "Иван"
    assert rows[0].patronymic == "Петрович"
    assert rows[0].row_number == 2


def test_read_excel_skips_empty_rows(tmp_path: Path):
    path = tmp_path / "empty_rows.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["Иванов", "Иван", "Петрович"])
    ws.append([None, None, None])  # пустая строка
    ws.append(["Петров", "Пётр", "Сергеевич"])
    wb.save(path)
    wb.close()

    config = AppConfig(excel_path=path, db_path=Path("unused.db"), start_row=2)
    rows = read_excel(config)
    assert len(rows) == 2


def test_read_excel_strips_whitespace(tmp_path: Path):
    path = tmp_path / "whitespace.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Фамилия", "Имя", "Отчество"])
    ws.append(["  Иванов  ", " Иван ", " Петрович "])
    wb.save(path)
    wb.close()

    config = AppConfig(excel_path=path, db_path=Path("unused.db"), start_row=2)
    rows = read_excel(config)
    assert rows[0].last_name == "Иванов"
    assert rows[0].first_name == "Иван"
    assert rows[0].patronymic == "Петрович"


def test_read_excel_custom_columns(tmp_path: Path):
    path = tmp_path / "custom_cols.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["ID", "Фамилия", "Имя", "Отчество"])
    ws.append([1, "Иванов", "Иван", "Петрович"])
    wb.save(path)
    wb.close()

    config = AppConfig(
        excel_path=path, db_path=Path("unused.db"), last_name_column="B", first_name_column="C", patronymic_column="D",
        start_row=2,
    )
    rows = read_excel(config)
    assert rows[0].last_name == "Иванов"


def test_read_excel_custom_start_row(tmp_path: Path):
    path = tmp_path / "custom_start.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Заголовок 1"])
    ws.append(["Заголовок 2"])
    ws.append(["Иванов", "Иван", "Петрович"])
    wb.save(path)
    wb.close()

    config = AppConfig(excel_path=path, db_path=Path("unused.db"), start_row=3)
    rows = read_excel(config)
    assert len(rows) == 1
    assert rows[0].row_number == 3


def test_read_excel_no_header(tmp_path: Path):
    """Тест: файл без заголовков, start_row=1 по умолчанию."""
    path = tmp_path / "no_header.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Иванов", "Иван", "Иванович"])
    ws.append(["Петров", "Пётр", "Петрович"])
    wb.save(path)
    wb.close()

    config = AppConfig(excel_path=path, db_path=Path("unused.db"))
    rows = read_excel(config)
    assert len(rows) == 2
    assert rows[0].row_number == 1
    assert rows[0].last_name == "Иванов"
    assert rows[1].row_number == 2
    assert rows[1].last_name == "Петров"


def test_get_output_path():
    assert get_output_path(Path("data.xlsx")) == Path("data_checked.xlsx")
    assert get_output_path(Path("/home/user/file.xlsx")) == Path("/home/user/file_checked.xlsx")
