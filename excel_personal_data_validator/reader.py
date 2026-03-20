import dataclasses

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from excel_personal_data_validator.config import AppConfig


@dataclasses.dataclass(frozen=True)
class PersonRow:
    """Одна строка с данными ФИО из Excel-файла."""

    row_number: int
    last_name: str
    first_name: str
    patronymic: str


def read_excel(config: AppConfig) -> list[PersonRow]:
    """Читает данные ФИО из Excel-файла.

    Использует openpyxl в режиме read_only для производительности.
    Пропускает строки, где все три поля пустые.
    """
    wb = load_workbook(config.excel_path, read_only=True, data_only=True)

    if config.sheet_name:
        ws = wb[config.sheet_name]
    else:
        ws = wb.active
        if ws is None:
            wb.close()
            msg = f"В файле '{config.excel_path}' нет активного листа"
            raise ValueError(msg)

    last_name_idx = column_index_from_string(config.last_name_column)
    first_name_idx = column_index_from_string(config.first_name_column)
    patronymic_idx = column_index_from_string(config.patronymic_column)

    rows: list[PersonRow] = []

    for row_number, row in enumerate(ws.iter_rows(min_row=config.start_row), start=config.start_row):
        last_name = _cell_value(row, last_name_idx)
        first_name = _cell_value(row, first_name_idx)
        patronymic = _cell_value(row, patronymic_idx)

        if not last_name and not first_name and not patronymic:
            continue

        rows.append(PersonRow(row_number=row_number, last_name=last_name, first_name=first_name, patronymic=patronymic))

    wb.close()
    return rows


def _cell_value(row: tuple[object, ...], col_index: int) -> str:
    """Извлекает строковое значение ячейки по индексу столбца (1-based)."""
    idx = col_index - 1
    if idx >= len(row):
        return ""
    cell = row[idx]
    value = cell.value if hasattr(cell, "value") else cell
    if value is None:
        return ""
    return str(value).strip()


def get_output_path(source_path: Path) -> Path:
    """Генерирует путь для файла с результатами проверки.

    data.xlsx -> data_checked.xlsx
    """
    return source_path.with_stem(f"{source_path.stem}_checked")
