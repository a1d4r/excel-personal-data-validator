from pathlib import Path

from openpyxl import load_workbook

from excel_personal_data_validator.config import AppConfig
from excel_personal_data_validator.db import NameCategory
from excel_personal_data_validator.models import ReviewSummary, UserAction


def category_to_column(category: NameCategory, config: AppConfig) -> str:
    """Возвращает букву столбца для указанной категории."""
    mapping = {
        NameCategory.LAST_NAME: config.last_name_column,
        NameCategory.FIRST_NAME: config.first_name_column,
        NameCategory.PATRONYMIC: config.patronymic_column,
    }
    return mapping[category]


def collect_corrections(summary: ReviewSummary, config: AppConfig) -> dict[tuple[int, str], str]:
    """Собирает исправления из решений пользователя."""
    corrections: dict[tuple[int, str], str] = {}
    for decision in summary.decisions:
        if decision.action == UserAction.REPLACE and decision.replacement:
            column = category_to_column(decision.entry.category, config)
            for row_num in decision.entry.row_numbers:
                corrections[(row_num, column)] = decision.replacement
    return corrections


def save_corrected_excel(
    source_path: Path, output_path: Path, corrections: dict[tuple[int, str], str], sheet_name: str | None = None
) -> None:
    """Сохраняет исправленную копию Excel-файла.

    Args:
        source_path: путь к исходному файлу.
        output_path: путь для сохранения копии.
        corrections: словарь (номер_строки, буква_столбца) -> новое_значение.
        sheet_name: имя листа (None = активный лист).
    """
    if not corrections:
        return

    wb = load_workbook(source_path)

    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active
        if ws is None:
            wb.close()
            msg = f"В файле '{source_path}' нет активного листа"
            raise ValueError(msg)

    for (row_number, column_letter), new_value in corrections.items():
        cell_ref = f"{column_letter}{row_number}"
        ws[cell_ref] = new_value

    wb.save(output_path)
    wb.close()
